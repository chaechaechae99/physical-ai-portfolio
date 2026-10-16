"""슬라이드 레이아웃. 각 함수는 (deck, spec, ctx) 를 받아 슬라이드를 추가한다.

ctx: theme, course, lecture, asset(kind, id) -> Path|None
비전공자용 원칙: 글은 최대 3줄, 그림은 크게.
"""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SW, SH = 13.333, 7.5
M = 0.6                     # 바깥 여백
BODY_TOP = 1.55             # 제목 아래 본문 시작


# ---------- 기본 도구 ----------

def rgb(ctx, name):
    return RGBColor.from_string(ctx["theme"]["colors"][name])


def _font(run, ctx, size, color, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(ctx, color)
    family = ctx["theme"]["font"]["family"]
    f.name = family
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):          # 한글(동아시아) 폰트도 지정해야 맑은 고딕으로 표시됨
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", family)


def text(slide, ctx, x, y, w, h, s, size=24, color="text", bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.05))
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        _font(r, ctx, size, color, bold)
    return tb


def box(slide, ctx, x, y, w, h, fill="light", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(ctx, fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def picture(slide, ctx, kind, ref, x, y, w, h, labels=None, cover=False):
    """이미지를 상자 안에 넣는다 (둥근 모서리 카드).

    cover=False: 비율 유지로 상자 안에 맞춤 / cover=True: 상자를 꽉 채우고 넘치는 쪽을 가운데 기준으로 자름.
    파일이 없으면 '대기 중' 자리표시 상자.
    """
    path = ctx["asset"](kind, ref)
    if path is None:
        b = box(slide, ctx, x, y, w, h, fill="light")
        b.line.color.rgb = rgb(ctx, "muted")
        b.line.dash_style = 4
        text(slide, ctx, x, y, w, h, f"[{kind} 대기 중]\n{ref}", 18, "muted",
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return (x, y, w, h)
    iw, ih = Image.open(path).size
    cx = cy = 0.0                       # 좌우·상하 각각 잘라낼 비율
    if cover:
        if iw / ih > w / h:
            cx = (1 - (w / h) / (iw / ih)) / 2
        else:
            cy = (1 - (iw / ih) / (w / h)) / 2
        pw, ph = w, h
    else:
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.crop_left = pic.crop_right = cx
    pic.crop_top = pic.crop_bottom = cy
    _round_corners(pic, 0.06)
    for lb in labels or []:
        lx = (lb["x"] - cx) / (1 - 2 * cx)          # 원본 이미지 기준 위치 → 잘린 이미지 기준
        ly = (lb["y"] - cy) / (1 - 2 * cy)
        _label(slide, ctx, lb["text"], lb.get("color", "text"), px + lx * pw, py + ly * ph)
    return (px, py, pw, ph)


def _round_corners(pic, radius):
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {int(radius * 100000)}")


def _label(slide, ctx, s, color, cx, cy):
    w = 0.32 * len(s) + 0.5
    box(slide, ctx, cx - w / 2, cy - 0.27, w, 0.54, fill=color, radius=0.5)
    text(slide, ctx, cx - w / 2, cy - 0.27, w, 0.54, s, 18, "bg", True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def blank(deck):
    return deck.slides.add_slide(deck.slide_layouts[6])


def header(slide, ctx, title):
    box(slide, ctx, M, 0.55, 0.12, 0.62, fill="primary", shape=MSO_SHAPE.RECTANGLE)
    text(slide, ctx, M + 0.3, 0.42, SW - 2 * M - 0.3, 0.9, title,
         ctx["theme"]["font"]["title_pt"], "text", True, anchor=MSO_ANCHOR.MIDDLE)


def footer(slide, ctx):
    lec = ctx["lecture"]
    small = ctx["theme"]["font"]["small_pt"]
    text(slide, ctx, M, SH - 0.5, 6, 0.35, ctx["course"]["course_title"], small, "muted")
    n = len(ctx["deck"].slides)
    text(slide, ctx, SW - M - 3, SH - 0.5, 3, 0.35, f"{lec['number']}강 · {n}", small, "muted",
         align=PP_ALIGN.RIGHT)


def notes(slide, spec):
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = spec["notes"]


def caption(slide, ctx, s, y):
    text(slide, ctx, M, y, SW - 2 * M, 0.6, s, ctx["theme"]["font"]["caption_pt"], "text", True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _media(spec):
    for kind in ("image", "video", "chart"):
        if kind in spec:
            return kind, spec[kind]
    return None, None


# ---------- 레이아웃 ----------

def title(deck, spec, ctx):
    s = blank(deck)
    lec = ctx["lecture"]
    box(s, ctx, SW * 0.55, 0, SW * 0.45, SH, fill="light", shape=MSO_SHAPE.RECTANGLE)
    tag = box(s, ctx, M, 2.0, 1.3, 0.6, fill="primary", radius=0.5)
    text(s, ctx, M, 2.0, 1.3, 0.6, f"{lec['number']}강", 22, "bg", True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, ctx, M, 2.8, 6.5, 1.4, lec["title"], 54, "text", True)
    text(s, ctx, M, 4.2, 6.5, 1.0, spec.get("subtitle", ""), 22, "muted")
    text(s, ctx, M, SH - 0.9, 6.5, 0.4, ctx["course"]["course_title"], 16, "muted")
    if "image" in spec:
        picture(s, ctx, "image", spec["image"], SW * 0.55 + 0.4, 0.6, SW * 0.45 - 0.8, SH - 1.2)
    notes(s, spec)


def goals(deck, spec, ctx):
    s = blank(deck)
    header(s, ctx, spec.get("title", "오늘 배울 것"))
    items = spec["items"][:3]
    gap = 0.4
    cw = (SW - 2 * M - gap * (len(items) - 1)) / len(items)
    for i, it in enumerate(items):
        x = M + i * (cw + gap)
        box(s, ctx, x, 2.0, cw, 4.2, fill="light")
        box(s, ctx, x + cw / 2 - 0.55, 2.5, 1.1, 1.1, fill="primary", shape=MSO_SHAPE.OVAL)
        text(s, ctx, x + cw / 2 - 0.55, 2.5, 1.1, 1.1, str(i + 1), 40, "bg", True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, ctx, x + 0.3, 3.9, cw - 0.6, 2.0, it, 26, "text", True, align=PP_ALIGN.CENTER)
    footer(s, ctx)
    notes(s, spec)


def big_figure(deck, spec, ctx):
    s = blank(deck)
    header(s, ctx, spec["title"])
    kind, ref = _media(spec)
    picture(s, ctx, kind, ref, M, BODY_TOP, SW - 2 * M, 4.75, spec.get("labels"))
    if spec.get("caption"):
        caption(s, ctx, spec["caption"], 6.35)
    footer(s, ctx)
    notes(s, spec)


def figure_text(deck, spec, ctx):
    s = blank(deck)
    header(s, ctx, spec["title"])
    kind, ref = _media(spec)
    picture(s, ctx, kind, ref, M, BODY_TOP, 6.6, 5.3, spec.get("labels"), cover=True)
    bullets = spec.get("bullets", [])[:3]
    x, w = 7.6, SW - M - 7.6
    for i, b in enumerate(bullets):
        y = BODY_TOP + 0.35 + i * 1.65
        box(s, ctx, x, y, w, 1.4, fill="light")
        box(s, ctx, x + 0.3, y + 0.55, 0.3, 0.3, fill="accent", shape=MSO_SHAPE.OVAL)
        text(s, ctx, x + 0.8, y, w - 1.0, 1.4, b, 22, "text", True, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, ctx)
    notes(s, spec)


def video(deck, spec, ctx):
    s = blank(deck)
    header(s, ctx, spec["title"])
    h = 4.7
    w = h * 16 / 9
    x, y = (SW - w) / 2, BODY_TOP
    mp4 = ctx["asset"]("video", spec["video"])
    if mp4 is None:
        picture(s, ctx, "video", spec["video"], x, y, w, h)
    else:
        poster = mp4.with_name(mp4.stem + "_poster.png")
        s.shapes.add_movie(str(mp4), Inches(x), Inches(y), Inches(w), Inches(h),
                           poster_frame_image=str(poster) if poster.exists() else None,
                           mime_type="video/mp4")
    if spec.get("caption"):
        caption(s, ctx, spec["caption"], 6.35)
    footer(s, ctx)
    notes(s, spec)


def quiz(deck, spec, ctx):
    """문제 슬라이드 + 정답 슬라이드 2장을 만든다."""
    for reveal in (False, True):
        s = blank(deck)
        header(s, ctx, spec.get("title", "퀴즈!") + (" — 정답" if reveal else ""))
        text(s, ctx, M, 1.8, SW - 2 * M, 1.0, spec["question"], 34, "text", True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        choices = spec["choices"]
        gap = 0.5
        cw = (SW - 2 * M - gap * (len(choices) - 1)) / len(choices)
        for i, c in enumerate(choices):
            x = M + i * (cw + gap)
            right = reveal and (i + 1) == spec["answer"]
            dim = reveal and not right
            box(s, ctx, x, 3.3, cw, 2.6, fill="success" if right else "light")
            text(s, ctx, x, 3.45, cw, 0.7, "①②③④⑤"[i], 30, "bg" if right else "muted", True,
                 align=PP_ALIGN.CENTER)
            text(s, ctx, x, 4.1, cw, 1.4, c, 32, "bg" if right else ("muted" if dim else "text"), True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        footer(s, ctx)
        notes(s, spec)


def summary(deck, spec, ctx):
    s = blank(deck)
    header(s, ctx, spec.get("title", "오늘의 정리"))
    has_img = "image" in spec
    w = 8.0 if has_img else SW - 2 * M
    for i, p in enumerate(spec["points"][:3]):
        y = BODY_TOP + 0.3 + i * 1.6
        box(s, ctx, M, y, w, 1.35, fill="light")
        box(s, ctx, M + 0.3, y + 0.3, 0.75, 0.75, fill="primary", shape=MSO_SHAPE.OVAL)
        text(s, ctx, M + 0.3, y + 0.3, 0.75, 0.75, str(i + 1), 26, "bg", True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, ctx, M + 1.35, y, w - 1.6, 1.35, p, 26, "text", True, anchor=MSO_ANCHOR.MIDDLE)
    if has_img:
        picture(s, ctx, "image", spec["image"], M + w + 0.3, BODY_TOP, SW - 2 * M - w - 0.3, 5.0)
    footer(s, ctx)
    notes(s, spec)


LAYOUTS = {f.__name__: f for f in (title, goals, big_figure, figure_text, video, quiz, summary)}
